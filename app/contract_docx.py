from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt
from html2docx import html2docx
from lxml import etree

APP_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = APP_ROOT.parent if (APP_ROOT.parent / "data").exists() else APP_ROOT
CONTRACTS_DIR = PROJECT_ROOT / "data" / "contracts"
HTML_BODY_RE = re.compile(r"<body.*?>.*?</body>", re.IGNORECASE | re.DOTALL)
SPACE_RE = re.compile(r"[ \t]+")


class _HtmlTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"p", "tr", "table", "body"}:
            self.parts.append("\n")
        elif tag == "td":
            self.parts.append("\t")

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"p", "tr", "table"}:
            self.parts.append("\n")

    def get_text(self) -> str:
        text = "".join(self.parts)
        lines = [SPACE_RE.sub(" ", line).strip() for line in text.splitlines()]
        return "\n".join(line for line in lines if line)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    heading = styles["Heading 1"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(14)
    heading.font.bold = True


def _normalize_contract_html(contract_html: str) -> str:
    html = contract_html.strip()
    match = HTML_BODY_RE.search(html)
    if match:
        return match.group(0)
    if html.startswith("<"):
        return f"<body>{html}</body>"
    escaped = html.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return f"<body><p>{escaped}</p></body>"


def contract_html_to_text(contract_html: str) -> str:
    parser = _HtmlTextExtractor()
    parser.feed(_normalize_contract_html(contract_html))
    parser.close()
    return parser.get_text()


def _contract_heading(deal_type: str | None) -> str:
    raw = (deal_type or "").strip()
    if raw.lower().startswith("договор"):
        raw = raw[len("договор"):].strip()
    suffix = raw.upper() if raw else "________________"
    return f"ДОГОВОР {suffix}"


def _build_header(document: Document, deal_type: str | None) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(_contract_heading(deal_type) + "\n")
    run.bold = True
    run.font.size = Pt(14)
    run2 = p.add_run("№ ______")
    run2.font.size = Pt(14)

    table = document.add_table(rows=1, cols=2)
    cell_left = table.rows[0].cells[0]
    cell_right = table.rows[0].cells[1]
    cell_left.text = "г. ___________"
    p_right = cell_right.paragraphs[0]
    p_right.text = "«___» _____________ 20___ г."
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph("")


def _apply_body_formatting(document: Document) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text.upper().startswith("ПРИЛОЖЕНИЕ №"):
            paragraph.paragraph_format.page_break_before = True
        if paragraph.alignment is None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph_format = paragraph.paragraph_format
        paragraph_format.first_line_indent = Cm(1.25)
        paragraph_format.space_after = Pt(0)
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)


def generate_docx(contract_html: str, case_id: str, version: int, deal_type: str | None = None) -> str:
    normalized_html = _normalize_contract_html(contract_html)
    buffer = html2docx(normalized_html, title="Договор")
    html_document = Document(buffer)

    final_document = Document()
    _configure_document(final_document)
    _build_header(final_document, deal_type)

    body_html = html_document.element.body
    body_final = final_document.element.body

    for child in body_html.iterchildren():
        cloned = etree.fromstring(etree.tostring(child))
        body_final.append(cloned)

    _apply_body_formatting(final_document)

    output_dir = CONTRACTS_DIR / str(case_id)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"v{version}.docx"
    final_document.save(output_path)
    return str(output_path)
